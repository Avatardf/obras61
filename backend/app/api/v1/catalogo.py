"""Catálogo de materiais da construtora.

Cada construtora tem o próprio catálogo. No primeiro acesso, a base padrão
do sistema (Excel com ~5 mil itens) é copiada para ela. Depois disso a
construtora inclui, edita, oculta, exclui e importa planilhas à vontade.
"""
import csv
import io
import unicodedata
import uuid
from typing import Annotated

import openpyxl
from fastapi import APIRouter, Depends, File, Form, HTTPException, Query, UploadFile, status
from fastapi.responses import StreamingResponse
from pydantic import BaseModel, Field
from sqlalchemy import func, insert, or_, select
from sqlalchemy.ext.asyncio import AsyncSession

from app.database import get_db
from app.dependencies import CurrentUser, require_papel
from app.models.catalogo import MaterialCatalogo
from app.models.tenant import Papel, Tenant, User
from app.utils.catalogo import carregar_catalogo

router = APIRouter(tags=["catalogo"])
DB = Depends(get_db)
# Quem pode alterar o catálogo
EditorCatalogo = Annotated[User, Depends(require_papel(Papel.admin, Papel.engenheiro, Papel.comprador))]

MAX_ARQUIVO = 10 * 1024 * 1024   # 10 MB


# ── Schemas ────────────────────────────────────────────────────────────────────

class MaterialOut(BaseModel):
    id: uuid.UUID
    codigo: str | None
    descricao: str
    unidade: str
    familia: str | None
    preco_referencia: float | None
    ativo: bool
    origem: str

    model_config = {"from_attributes": True}


class MaterialIn(BaseModel):
    codigo: str | None = Field(default=None, max_length=50)
    descricao: str = Field(min_length=2, max_length=400)
    unidade: str = Field(default="un", min_length=1, max_length=20)
    familia: str | None = Field(default=None, max_length=120)
    preco_referencia: float | None = Field(default=None, ge=0)


class MaterialPatch(BaseModel):
    codigo: str | None = Field(default=None, max_length=50)
    descricao: str | None = Field(default=None, min_length=2, max_length=400)
    unidade: str | None = Field(default=None, min_length=1, max_length=20)
    familia: str | None = Field(default=None, max_length=120)
    preco_referencia: float | None = Field(default=None, ge=0)
    ativo: bool | None = None


# ── Helpers ────────────────────────────────────────────────────────────────────

def _norm(texto) -> str:
    """Normaliza para comparar: sem acento, maiúsculas, espaços simples."""
    if texto is None:
        return ""
    s = unicodedata.normalize("NFKD", str(texto)).encode("ascii", "ignore").decode()
    return " ".join(s.upper().split())


def _limpo(v, tam: int) -> str | None:
    if v is None:
        return None
    s = " ".join(str(v).split())
    return s[:tam] if s else None


async def garantir_catalogo(db: AsyncSession, tenant_id: uuid.UUID) -> None:
    """Copia a base padrão para a construtora no primeiro acesso (uma vez só)."""
    pronto = (await db.execute(
        select(Tenant.catalogo_inicializado).where(Tenant.id == tenant_id)
    )).scalar_one()
    if pronto:
        return
    # Trava a linha do tenant: duas requisições simultâneas não copiam em dobro
    tenant = await db.get(Tenant, tenant_id, with_for_update=True, populate_existing=True)
    if tenant.catalogo_inicializado:
        await db.commit()
        return
    linhas = [
        {
            "id": uuid.uuid4(), "tenant_id": tenant_id,
            "codigo": _limpo(m["codigo"], 50), "descricao": _limpo(m["descricao"], 400),
            "unidade": _limpo(m["unidade"], 20) or "un", "familia": _limpo(m["familia"], 120),
            "ativo": True, "origem": "padrao",
        }
        for m in carregar_catalogo() if _limpo(m["descricao"], 400)
    ]
    if linhas:
        await db.execute(insert(MaterialCatalogo), linhas)
    tenant.catalogo_inicializado = True
    await db.commit()


async def _material(db: AsyncSession, mid: uuid.UUID, tenant_id) -> MaterialCatalogo:
    m = await db.get(MaterialCatalogo, mid)
    if not m or m.tenant_id != tenant_id:
        raise HTTPException(status.HTTP_404_NOT_FOUND, "Material não encontrado")
    return m


async def _descricao_em_uso(db: AsyncSession, tenant_id, descricao: str, exceto: uuid.UUID | None = None) -> bool:
    stmt = select(MaterialCatalogo.id).where(
        MaterialCatalogo.tenant_id == tenant_id,
        func.lower(MaterialCatalogo.descricao) == descricao.strip().lower(),
    )
    if exceto:
        stmt = stmt.where(MaterialCatalogo.id != exceto)
    return (await db.execute(stmt.limit(1))).first() is not None


# ── Leitura ────────────────────────────────────────────────────────────────────

@router.get("/materiais")
async def listar_materiais(
    q: str | None = Query(None, description="Busca por descrição ou código (todas as palavras)"),
    familia: str | None = Query(None, description="Filtrar por família exata"),
    limit: int = Query(50, ge=1, le=6000),
    offset: int = Query(0, ge=0),
    incluir_inativos: bool = Query(False, description="Inclui os itens ocultos"),
    db: AsyncSession = DB, user: CurrentUser = None,
):
    """Materiais da construtora, com busca, filtro por família e paginação."""
    await garantir_catalogo(db, user.tenant_id)
    stmt = select(MaterialCatalogo).where(MaterialCatalogo.tenant_id == user.tenant_id)
    if not incluir_inativos:
        stmt = stmt.where(MaterialCatalogo.ativo == True)  # noqa: E712
    if familia:
        stmt = stmt.where(MaterialCatalogo.familia == familia)
    if q:
        for palavra in q.split():
            termo = f"%{palavra}%"
            stmt = stmt.where(or_(MaterialCatalogo.descricao.ilike(termo), MaterialCatalogo.codigo.ilike(termo)))

    total = (await db.execute(select(func.count()).select_from(stmt.subquery()))).scalar_one()
    itens = (await db.execute(
        stmt.order_by(MaterialCatalogo.descricao).offset(offset).limit(limit)
    )).scalars().all()
    return {"total": total, "items": [MaterialOut.model_validate(m) for m in itens]}


@router.get("/materiais/familias")
async def listar_familias_endpoint(db: AsyncSession = DB, user: CurrentUser = None):
    """Famílias em uso no catálogo da construtora."""
    await garantir_catalogo(db, user.tenant_id)
    rows = (await db.execute(
        select(MaterialCatalogo.familia).where(
            MaterialCatalogo.tenant_id == user.tenant_id,
            MaterialCatalogo.ativo == True,  # noqa: E712
            MaterialCatalogo.familia.is_not(None),
        ).distinct().order_by(MaterialCatalogo.familia)
    )).scalars().all()
    return list(rows)


@router.get("/materiais/modelo")
async def baixar_modelo(user: CurrentUser = None):
    """Planilha modelo para importação."""
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Materiais"
    ws.append(["Código", "Descrição", "Unidade", "Família", "Preço"])
    ws.append(["CIM-01", "Cimento CP II-E 50 kg", "SC", "ESTRUTURA", 38.9])
    ws.append(["", "Areia média lavada", "M3", "AGREGADOS", 145])
    for col, largura in zip("ABCDE", (12, 48, 10, 20, 12)):
        ws.column_dimensions[col].width = largura
    buf = io.BytesIO()
    wb.save(buf)
    buf.seek(0)
    return StreamingResponse(
        buf, media_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
        headers={"Content-Disposition": 'attachment; filename="modelo_catalogo_materiais.xlsx"'},
    )


# ── Escrita ────────────────────────────────────────────────────────────────────

@router.post("/materiais", response_model=MaterialOut, status_code=status.HTTP_201_CREATED)
async def criar_material(body: MaterialIn, db: AsyncSession = DB, user: EditorCatalogo = None):
    await garantir_catalogo(db, user.tenant_id)
    if await _descricao_em_uso(db, user.tenant_id, body.descricao):
        raise HTTPException(status.HTTP_409_CONFLICT, "Já existe um material com esta descrição.")
    m = MaterialCatalogo(
        id=uuid.uuid4(), tenant_id=user.tenant_id, origem="proprio", ativo=True,
        codigo=_limpo(body.codigo, 50), descricao=_limpo(body.descricao, 400),
        unidade=_limpo(body.unidade, 20) or "un", familia=_limpo(body.familia, 120),
        preco_referencia=body.preco_referencia,
    )
    db.add(m)
    await db.commit()
    await db.refresh(m)
    return m


@router.patch("/materiais/{mid}", response_model=MaterialOut)
async def atualizar_material(mid: uuid.UUID, body: MaterialPatch, db: AsyncSession = DB, user: EditorCatalogo = None):
    m = await _material(db, mid, user.tenant_id)
    dados = body.model_dump(exclude_unset=True)
    if dados.get("descricao") and await _descricao_em_uso(db, user.tenant_id, dados["descricao"], exceto=m.id):
        raise HTTPException(status.HTTP_409_CONFLICT, "Já existe um material com esta descrição.")
    tamanhos = {"codigo": 50, "descricao": 400, "unidade": 20, "familia": 120}
    for campo, valor in dados.items():
        if campo in tamanhos:
            valor = _limpo(valor, tamanhos[campo])
            if campo in ("descricao", "unidade") and not valor:
                continue   # campos obrigatórios não ficam vazios
        setattr(m, campo, valor)
    await db.commit()
    await db.refresh(m)
    return m


@router.delete("/materiais/{mid}", status_code=status.HTTP_204_NO_CONTENT)
async def excluir_material(mid: uuid.UUID, db: AsyncSession = DB, user: EditorCatalogo = None):
    """Exclui de vez. Orçamentos e requisições antigos guardam a descrição em
    texto, então não são afetados."""
    m = await _material(db, mid, user.tenant_id)
    await db.delete(m)
    await db.commit()


# ── Importação de planilha ─────────────────────────────────────────────────────

COLUNAS = {
    "codigo":    {"CODIGO", "COD", "COD.", "CODIGO DO MATERIAL", "REFERENCIA", "REF"},
    "descricao": {"DESCRICAO", "DESCRICAO DO MATERIAL", "MATERIAL", "ITEM", "NOME", "PRODUTO", "DESCRICÃO"},
    "unidade":   {"UNIDADE", "UND", "UN", "UNID", "UNID.", "UN."},
    "familia":   {"FAMILIA", "GRUPO", "CATEGORIA", "CLASSE"},
    "preco":     {"PRECO", "PRECO UNITARIO", "VALOR", "VALOR UNITARIO", "CUSTO", "CUSTO UNITARIO", "PRECO REFERENCIA"},
}


def _preco(v) -> float | None:
    """Aceita número ou texto pt-BR/en-US ('1.234,56', 'R$ 38,90', '38.9')."""
    if v is None or v == "":
        return None
    if isinstance(v, (int, float)):
        return float(v) if v >= 0 else None
    s = str(v).replace("R$", "").strip()
    if "," in s:
        s = s.replace(".", "").replace(",", ".")
    try:
        n = float(s)
        return n if n >= 0 else None
    except ValueError:
        raise ValueError(f"preço inválido '{v}'")


def _ler_linhas(nome: str, conteudo: bytes) -> list[list]:
    if nome.lower().endswith(".csv"):
        try:
            texto = conteudo.decode("utf-8-sig")
        except UnicodeDecodeError:
            texto = conteudo.decode("latin-1")
        delim = ";" if texto.split("\n", 1)[0].count(";") >= texto.split("\n", 1)[0].count(",") else ","
        return [linha for linha in csv.reader(io.StringIO(texto), delimiter=delim)]
    try:
        wb = openpyxl.load_workbook(io.BytesIO(conteudo), read_only=True, data_only=True)
    except Exception:
        raise HTTPException(status.HTTP_422_UNPROCESSABLE_ENTITY, "Arquivo inválido. Envie uma planilha .xlsx ou .csv.")
    linhas = [list(r) for r in wb.active.iter_rows(values_only=True)]
    wb.close()
    return linhas


def _mapear_cabecalho(linhas: list[list]) -> tuple[int, dict[str, int]]:
    """Procura o cabeçalho nas 10 primeiras linhas; exige a coluna de descrição."""
    for i, linha in enumerate(linhas[:10]):
        mapa: dict[str, int] = {}
        for j, cel in enumerate(linha):
            n = _norm(cel)
            for campo, nomes in COLUNAS.items():
                if n in {_norm(x) for x in nomes} and campo not in mapa:
                    mapa[campo] = j
        if "descricao" in mapa:
            return i, mapa
    raise HTTPException(
        status.HTTP_422_UNPROCESSABLE_ENTITY,
        "Não encontrei a coluna 'Descrição'. Use o modelo: Código, Descrição, Unidade, Família, Preço.",
    )


@router.post("/materiais/importar")
async def importar_planilha(
    arquivo: UploadFile = File(...),
    modo: str = Form("adicionar"),        # adicionar | substituir
    simular: bool = Form(True),           # True = só calcula, não grava
    db: AsyncSession = DB, user: EditorCatalogo = None,
):
    """Importa materiais de uma planilha.

    - adicionar: cria os novos e atualiza os existentes (casados por código
      ou, sem código, pela descrição). Nada é removido.
    - substituir: igual ao adicionar e, além disso, OCULTA os itens que não
      estão na planilha. Ocultar é reversível; nada é apagado.
    - simular: devolve o resumo sem gravar nada.
    """
    if modo not in ("adicionar", "substituir"):
        raise HTTPException(status.HTTP_422_UNPROCESSABLE_ENTITY, "Modo inválido.")
    conteudo = await arquivo.read()
    if len(conteudo) > MAX_ARQUIVO:
        raise HTTPException(status.HTTP_413_REQUEST_ENTITY_TOO_LARGE, "Arquivo maior que 10 MB.")

    linhas = _ler_linhas(arquivo.filename or "", conteudo)
    i_cab, mapa = _mapear_cabecalho(linhas)

    await garantir_catalogo(db, user.tenant_id)
    existentes = (await db.execute(
        select(MaterialCatalogo).where(MaterialCatalogo.tenant_id == user.tenant_id)
    )).scalars().all()
    por_codigo = {_norm(m.codigo): m for m in existentes if m.codigo}
    por_desc = {_norm(m.descricao): m for m in existentes}

    def cel(linha, campo):
        j = mapa.get(campo)
        return linha[j] if j is not None and j < len(linha) else None

    novos = atualizados = sem_alteracao = 0
    erros: list[str] = []
    vistos: set[uuid.UUID] = set()

    for n_linha, linha in enumerate(linhas[i_cab + 1:], start=i_cab + 2):
        descricao = _limpo(cel(linha, "descricao"), 400)
        if not descricao:
            continue   # linha vazia
        try:
            preco = _preco(cel(linha, "preco"))
        except ValueError as e:
            erros.append(f"Linha {n_linha}: {e}")
            continue
        codigo = _limpo(cel(linha, "codigo"), 50)
        unidade = _limpo(cel(linha, "unidade"), 20)
        familia = _limpo(cel(linha, "familia"), 120)

        m = (por_codigo.get(_norm(codigo)) if codigo else None) or por_desc.get(_norm(descricao))
        if m:
            if m.id in vistos:
                erros.append(f"Linha {n_linha}: '{descricao}' repetido na planilha (ignorado)")
                continue
            vistos.add(m.id)
            # Só sobrescreve o que a planilha trouxe preenchido
            mudancas = {"descricao": descricao, "codigo": codigo, "unidade": unidade,
                        "familia": familia, "preco_referencia": preco, "ativo": True}
            alterou = False
            for campo, valor in mudancas.items():
                if valor is None:
                    continue
                atual = getattr(m, campo)
                if campo == "preco_referencia" and atual is not None:
                    atual = float(atual)
                if atual != valor:
                    setattr(m, campo, valor)
                    alterou = True
            if alterou:
                atualizados += 1
            else:
                sem_alteracao += 1
        else:
            m = MaterialCatalogo(
                id=uuid.uuid4(), tenant_id=user.tenant_id, origem="planilha", ativo=True,
                codigo=codigo, descricao=descricao, unidade=unidade or "un",
                familia=familia, preco_referencia=preco,
            )
            db.add(m)
            vistos.add(m.id)
            por_desc[_norm(descricao)] = m
            if codigo:
                por_codigo[_norm(codigo)] = m
            novos += 1

    ocultados = 0
    if modo == "substituir":
        for m in existentes:
            if m.id not in vistos and m.ativo:
                m.ativo = False
                ocultados += 1

    if simular:
        await db.rollback()
    else:
        await db.commit()

    return {
        "simulado": simular, "modo": modo,
        "novos": novos, "atualizados": atualizados, "sem_alteracao": sem_alteracao,
        "ocultados": ocultados, "erros": erros[:50], "total_erros": len(erros),
    }
